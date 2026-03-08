from __future__ import annotations

from dataclasses import dataclass
import os
from pathlib import Path
from typing import Any
import csv
import json

import chardet
import openpyxl
import pandas as pd
from docx import Document
from pypdf import PdfReader

from .inventory import ensure_inventory_schema


TEXT_LIKE_SUFFIXES = {
    ".txt",
    ".md",
    ".csv",
    ".tsv",
    ".json",
    ".yaml",
    ".yml",
    ".xml",
    ".html",
    ".htm",
    ".log",
    ".ini",
    ".cfg",
    ".conf",
    ".py",
    ".ps1",
    ".bat",
    ".sh",
}
PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}
XLSX_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}


def _to_os_path(path: str | Path) -> str:
    """Return an OS path string suitable for Win32 long-path access when needed."""
    raw = str(path)
    if os.name != "nt":
        return raw

    normalized = raw.replace("/", "\\")
    if normalized.startswith("\\\\?\\"):
        return normalized
    if normalized.startswith("\\\\"):
        return "\\\\?\\UNC\\" + normalized.lstrip("\\")
    return "\\\\?\\" + normalized


@dataclass(frozen=True)
class ExtractConfig:
    max_chars_per_file: int = 12000
    max_csv_rows: int = 30
    max_csv_columns: int = 20
    max_xlsx_rows_per_sheet: int = 30
    max_xlsx_columns: int = 20
    max_docx_paragraphs: int = 300
    max_pdf_pages: int = 30
    preview_chars: int = 300


@dataclass(frozen=True)
class ExtractResult:
    text_status: str
    text_source: str
    extracted_text: str
    text_preview: str
    extracted_chars: int
    text_truncated: bool
    text_error: str | None
    extracted_pages: int | None = None
    extracted_sheets: int | None = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "text_status": self.text_status,
            "text_source": self.text_source,
            "extracted_text": self.extracted_text,
            "text_preview": self.text_preview,
            "extracted_chars": self.extracted_chars,
            "text_truncated": self.text_truncated,
            "text_error": self.text_error,
            "extracted_pages": self.extracted_pages,
            "extracted_sheets": self.extracted_sheets,
        }


def _normalize_text(value: str) -> str:
    lines = [line.strip() for line in value.replace("\x00", " ").splitlines()]
    cleaned = "\n".join(line for line in lines if line)
    return cleaned.strip()


def _truncate_text(text: str, limit: int) -> tuple[str, bool]:
    if len(text) <= limit:
        return text, False
    return text[:limit].rstrip() + "\n...[TRUNCATED]", True


def _preview(text: str, limit: int) -> str:
    compact = " ".join(text.split())
    if len(compact) <= limit:
        return compact
    return compact[:limit].rstrip() + "..."


def _read_text_file(path: Path, config: ExtractConfig) -> ExtractResult:
    with open(_to_os_path(path), "rb") as f:
        raw = f.read()
    encoding_guess = chardet.detect(raw).get("encoding") or "utf-8"
    try:
        text = raw.decode(encoding_guess, errors="replace")
    except LookupError:
        text = raw.decode("utf-8", errors="replace")

    if path.suffix.lower() == ".csv":
        reader = csv.reader(text.splitlines())
        rows: list[list[str]] = []
        for i, row in enumerate(reader):
            if i >= config.max_csv_rows:
                break
            rows.append([str(cell) for cell in row[: config.max_csv_columns]])
        text = "\n".join(", ".join(cell.strip() for cell in row) for row in rows)
    elif path.suffix.lower() == ".json":
        try:
            parsed = json.loads(text)
            text = json.dumps(parsed, ensure_ascii=False, indent=2)
        except Exception:
            pass

    normalized = _normalize_text(text)
    extracted_text, truncated = _truncate_text(normalized, config.max_chars_per_file)
    return ExtractResult(
        text_status="ok",
        text_source="text",
        extracted_text=extracted_text,
        text_preview=_preview(extracted_text, config.preview_chars),
        extracted_chars=len(extracted_text),
        text_truncated=truncated,
        text_error=None,
    )


def _read_pdf(path: Path, config: ExtractConfig) -> ExtractResult:
    with open(_to_os_path(path), "rb") as f:
        reader = PdfReader(f)
        pages_text: list[str] = []
        pages_read = 0
        for page in reader.pages[: config.max_pdf_pages]:
            pages_read += 1
            pages_text.append(page.extract_text() or "")
    normalized = _normalize_text("\n\n".join(pages_text))
    extracted_text, truncated = _truncate_text(normalized, config.max_chars_per_file)
    status = "ok" if extracted_text else "empty"
    return ExtractResult(
        text_status=status,
        text_source="pdf",
        extracted_text=extracted_text,
        text_preview=_preview(extracted_text, config.preview_chars),
        extracted_chars=len(extracted_text),
        text_truncated=truncated,
        text_error=None,
        extracted_pages=pages_read,
    )


def _read_docx(path: Path, config: ExtractConfig) -> ExtractResult:
    with open(_to_os_path(path), "rb") as f:
        doc = Document(f)
    chunks: list[str] = []
    para_count = 0
    for para in doc.paragraphs:
        if para_count >= config.max_docx_paragraphs:
            break
        if para.text.strip():
            chunks.append(para.text)
        para_count += 1

    for table in doc.tables:
        for row in table.rows[: config.max_csv_rows]:
            row_text = [cell.text.strip() for cell in row.cells[: config.max_csv_columns]]
            if any(row_text):
                chunks.append(" | ".join(row_text))

    normalized = _normalize_text("\n".join(chunks))
    extracted_text, truncated = _truncate_text(normalized, config.max_chars_per_file)
    status = "ok" if extracted_text else "empty"
    return ExtractResult(
        text_status=status,
        text_source="docx",
        extracted_text=extracted_text,
        text_preview=_preview(extracted_text, config.preview_chars),
        extracted_chars=len(extracted_text),
        text_truncated=truncated,
        text_error=None,
    )


def _read_xlsx(path: Path, config: ExtractConfig) -> ExtractResult:
    wb = openpyxl.load_workbook(_to_os_path(path), read_only=True, data_only=True)
    parts: list[str] = []
    sheet_count = 0
    for ws in wb.worksheets:
        sheet_count += 1
        parts.append(f"[SHEET] {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= config.max_xlsx_rows_per_sheet:
                break
            values = []
            for value in list(row)[: config.max_xlsx_columns]:
                if value is None:
                    values.append("")
                else:
                    values.append(str(value).strip())
            if any(values):
                parts.append(" | ".join(values))
    normalized = _normalize_text("\n".join(parts))
    extracted_text, truncated = _truncate_text(normalized, config.max_chars_per_file)
    status = "ok" if extracted_text else "empty"
    return ExtractResult(
        text_status=status,
        text_source="xlsx",
        extracted_text=extracted_text,
        text_preview=_preview(extracted_text, config.preview_chars),
        extracted_chars=len(extracted_text),
        text_truncated=truncated,
        text_error=None,
        extracted_sheets=sheet_count,
    )


def extract_text_from_path(path: str | Path, config: ExtractConfig | None = None) -> ExtractResult:
    config = config or ExtractConfig()
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    try:
        if suffix in TEXT_LIKE_SUFFIXES:
            return _read_text_file(file_path, config)
        if suffix in PDF_SUFFIXES:
            return _read_pdf(file_path, config)
        if suffix in DOCX_SUFFIXES:
            return _read_docx(file_path, config)
        if suffix in XLSX_SUFFIXES:
            return _read_xlsx(file_path, config)
        return ExtractResult(
            text_status="unsupported",
            text_source="unsupported",
            extracted_text="",
            text_preview="",
            extracted_chars=0,
            text_truncated=False,
            text_error=None,
        )
    except Exception as exc:
        return ExtractResult(
            text_status="error",
            text_source=suffix.lstrip(".") or "unknown",
            extracted_text="",
            text_preview="",
            extracted_chars=0,
            text_truncated=False,
            text_error=f"{type(exc).__name__}: {exc}",
        )


def enrich_inventory_with_text(
    inventory_df: pd.DataFrame,
    path_column: str = "absolute_path",
    config: ExtractConfig | None = None,
) -> pd.DataFrame:
    config = config or ExtractConfig()
    if inventory_df.empty:
        return inventory_df.copy()
    inventory_df = ensure_inventory_schema(inventory_df)
    if path_column not in inventory_df.columns:
        raise KeyError(f"Missing path column: {path_column}")

    results = inventory_df[path_column].apply(lambda p: extract_text_from_path(p, config).as_dict())
    extracted = pd.DataFrame(list(results))
    out = pd.concat([inventory_df.reset_index(drop=True), extracted.reset_index(drop=True)], axis=1)
    out["has_extracted_text"] = out["extracted_chars"].fillna(0).astype(int) > 0
    return out


def save_text_outputs(df: pd.DataFrame, output_base: str | Path) -> tuple[Path, Path]:
    output_base = Path(output_base)
    output_base.parent.mkdir(parents=True, exist_ok=True)
    csv_path = output_base.with_suffix(".csv")
    parquet_path = output_base.with_suffix(".parquet")
    df.to_csv(csv_path, index=False, encoding="utf-8-sig")
    df.to_parquet(parquet_path, index=False)
    return csv_path, parquet_path
